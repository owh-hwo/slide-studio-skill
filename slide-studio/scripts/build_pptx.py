"""Convert an HTML slide deck to PPTX: one full-bleed picture per slide + speaker notes.

The HTML deck stays the source. Each slide is rendered by headless Chrome at
2x (3840x2160) so text stays sharp when projected, placed edge to edge on a
16:9 PowerPoint slide, and the slide's <aside class="notes"> goes into the
PowerPoint notes pane (bold kept), so Presenter View shows the speaker script.
Slide text is therefore not editable in PowerPoint: edit the HTML and re-run.

    python build_pptx.py "<deck>.html"

Writes the .pptx next to the deck. Needs python-pptx and Pillow (pip install python-pptx pillow).
PermissionError = the .pptx is open in PowerPoint (a ~$ lock file shows it); close it and re-run.
"""
import html
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image, ImageStat
from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, str(Path(__file__).resolve().parent))
from chrome import CHROME  # noqa: E402

W, H = 12192000, 6858000  # 13.333 x 7.5 in, 16:9


def render(deck, i, png):
    # a near-uniform image is a slide caught before it painted: retry, then fail loudly
    for _ in range(3):
        subprocess.run([CHROME, "--headless=new", "--hide-scrollbars", "--force-device-scale-factor=2",
                        "--window-size=1920,1080", "--virtual-time-budget=4000",
                        f"--screenshot={png}", deck.as_uri() + f"?preview={i}"],
                       capture_output=True, timeout=120)
        if png.exists() and ImageStat.Stat(Image.open(png).convert("L")).stddev[0] > 2:
            return
    raise SystemExit(f"render failed for slide {i} (missing or blank)")


def note_paragraphs(section):
    """<aside class="notes"> -> list of paragraphs, each a list of (text, bold) runs."""
    m = re.search(r'<aside class="notes">(.*?)</aside>', section, re.S)
    if not m:
        return []
    body = m.group(1)
    blocks = re.findall(r"<(?:p|li)\b[^>]*>(.*?)</(?:p|li)>", body, re.S) or [body]
    paras = []
    for b in blocks:
        runs = []
        for part in re.split(r"(<(?:strong|b)\b[^>]*>.*?</(?:strong|b)>)", b, flags=re.S):
            bold = bool(re.match(r"<(?:strong|b)\b", part))
            text = html.unescape(re.sub(r"<[^>]+>", "", part))
            text = re.sub(r"\s+", " ", text)
            if text.strip():
                runs.append((text, bold))
        if runs:
            runs[0] = (runs[0][0].lstrip(), runs[0][1])
            runs[-1] = (runs[-1][0].rstrip(), runs[-1][1])
            paras.append(runs)
    return paras


def main(deck_path):
    deck = Path(deck_path).resolve()
    src = deck.read_text(encoding="utf-8")
    sections = re.findall(r'<section class="slide\b.*?</section>', src, re.S)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    blank = prs.slide_layouts[6]
    with tempfile.TemporaryDirectory() as tmp:
        for i, sec in enumerate(sections, 1):
            png = Path(tmp) / f"s{i:02d}.png"
            render(deck, i, png)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(png), 0, 0, Emu(W), Emu(H))
            paras = note_paragraphs(sec)
            if paras:
                tf = slide.notes_slide.notes_text_frame
                for k, runs in enumerate(paras):
                    p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                    for text, bold in runs:
                        r = p.add_run()
                        r.text = text
                        r.font.bold = bold
            print(f"  slide {i}/{len(sections)}", end="\r")
    out = deck.with_suffix(".pptx")
    prs.save(out)
    print(f"{out.name}: {len(sections)} slides, {sum(1 for s in sections if 'class=\"notes\"' in s)} with notes")


if __name__ == "__main__":
    main(sys.argv[1])
